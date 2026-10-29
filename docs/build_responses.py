from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.75)
for n in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[n]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.06
doc.styles['Title'].font.size=Pt(24)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 2'].font.size=Pt(12)
header=sec.header.paragraphs[0]; header.text='MAD 4360  |  ACTIVITY 11B  |  OCTOBER 29 2026'; header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0]; footer.alignment=2
r=footer.add_run('Myles Miller and Eziz Bagshiyev  |  '); r.font.size=Pt(9)
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)

def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
 for c,w,h in zip(t.rows[0].cells,widths,headers): c.width=Inches(w); c.text=h
 for row in rows:
  cells=t.add_row().cells
  for c,w,val in zip(cells,widths,row): c.width=Inches(w); c.text=str(val)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trpr.append(cant)
  for c in row.cells:
   c.vertical_alignment=1
   pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   pr.append(borders)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); margins.append(e)
   pr.append(margins)
   shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E6ECF1' if ri==0 else ('F6F7F8' if ri%2==0 else 'FFFFFF')); pr.append(shade)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
    for r in p.runs: r.font.size=Pt(10); r.bold=ri==0
 return t

def image(name,width):
 p=doc.add_paragraph(); p.alignment=1; p.add_run().add_picture(str(root/'docs'/name),width=Inches(width))

def caption(text):
 p=doc.add_paragraph(text); p.paragraph_format.space_after=Pt(7)
 for r in p.runs: r.font.size=Pt(9); r.italic=True

doc.add_paragraph('Activity 11 Creature Registry Evidence',style='Title')
doc.add_paragraph('Myles Miller and Eziz Bagshiyev\nUndergraduate track   |   Shared build only')
doc.add_paragraph('Repository: https://github.com/ozemoya/incl11')
doc.add_paragraph('This document answers two Ghost Hunt prompts using the Nightfall Registry build, automated Flutter checks, and screenshots captured on the Pixel 7 Pro Android emulator. The tables link each observed result to its test or screenshot.')
doc.add_heading('The Poltergeist Field',level=1)
doc.add_paragraph('Rotating this build did not erase the name, visible password, birth date, or Black Cat selection. The name and password text live in controllers inside _SignupScreenState, while _obscurePassword, _birthDate, and _selectedCreatureType are State fields; _dateController mirrors the displayed date. Rotation changes the layout but keeps the existing State, so the values stay available. Leaving the signup route is different: dispose releases the controllers, and reopening the form creates empty fields with the password hidden again. If a value disappeared during an ordinary rebuild, the smallest fix would be to move its owner out of build and into State, while process-death recovery would require a separate design that avoids restoring passwords.')
table(['Value','Owner','Observed after rotation'],[
 ['Name and password','TextEditingControllers','Myles Miller and test input retained'],
 ['Password visibility','_obscurePassword','Still visible'],
 ['Birth date','_birthDate and _dateController','1/1/2008 retained'],
 ['Creature','_selectedCreatureType','Black Cat retained'],
],[1.3,2.8,2.9])
image('nightfall-landscape-state.png',6.5)
caption('Landscape capture from integration_test/registry_test.dart. The test also asserts the date and creature value even when they are below the viewport. The visible password is disposable test data.')

doc.add_heading('The Clairvoyant Ghost',level=1).paragraph_format.page_break_before=True
doc.add_paragraph('The five dates below check both sides of the proposed 1990 cutoff, a later year, the current date, and a future date. A picker ending in 1990 would incorrectly reject the 2008 and current-date rows, so testing only old birthdays would miss that bug. In this build, _pickBirthDate uses firstDate: DateTime(1920) and lastDate: DateUtils.dateOnly(DateTime.now()), and the Android run accepted all four allowed dates. Entering October 30, 2026 kept the picker open with “Out of range,” which matches the expected rejection of a future birthday. These results apply to the device date of October 29, 2026; the automated test calculates today and tomorrow again when rerun.')
table(['Input','Expected','Observed on emulator'],[
 ['12/31/1989','Accept','Accepted and displayed'],
 ['1/1/1990','Accept','Accepted and displayed'],
 ['1/1/2008','Accept','Accepted and displayed'],
 ['10/29/2026 today','Accept','Accepted and displayed'],
 ['10/30/2026 tomorrow','Reject','Out of range; dialog stayed open'],
],[1.5,1.1,4.4])
p=doc.add_paragraph(); p.alignment=1
p.add_run().add_picture(str(root/'docs/nightfall-date-2-accepted.png'),width=Inches(2.05))
p.add_run('    ')
p.add_run().add_picture(str(root/'docs/nightfall-date-future-rejected.png'),width=Inches(2.05))
caption('Left: a 2008 date accepted in the guest book. Right: tomorrow rejected by the date input. Other accepted-date captures are retained in docs/.')
doc.add_paragraph('Code trace: SignupScreen → _pickBirthDate → showDatePicker → mounted check → setState updates _birthDate and _dateController. Cancelling the picker returns without changing the previous selection.')

doc.add_heading('Flow and validation checks',level=1).paragraph_format.page_break_before=True
doc.add_paragraph('Verification on October 29 2026: flutter analyze found no issues; 14 widget tests and one Android integration test passed. The table distinguishes checks run in the widget suite from those also run on the emulator.')
table(['Check','Observed result','Evidence'],[
 ['Empty registration','All five required-field errors appeared','Widget and Android'],
 ['Email without @ and 5-character password','Both original validation messages appeared','Widget test'],
 ['All five creatures','Each reached a confirmation with its own type','Five widget cases'],
 ['Eye toggle','Show and hide changed password visibility','Widget and Android'],
 ['Date rules','Empty blocked; future rejected; valid dates kept','Widget and Android'],
 ['Back from signup','Foyer returned; new form was empty','Widget and Android'],
 ['Back from success','Foyer returned; signup was not underneath','Widget test'],
 ['Loading','Register button disabled; one success screen','Widget test'],
 ['Exit during loading','No success route or disposed-state exception','Widget test'],
 ['Confetti and return','Success opened; foyer and new form opened safely','Android integration'],
],[1.65,3.65,1.7])
p=doc.add_paragraph(); p.alignment=1
p.add_run().add_picture(str(root/'docs/nightfall-welcome.png'),width=Inches(1.5)); p.add_run('      ')
p.add_run().add_picture(str(root/'docs/nightfall-success.png'),width=Inches(1.5))
caption('The entry screen and personalized ballroom, captured from the running Android app. The form is replaced on success, so the previous guest book cannot be resubmitted with Back.')
doc.add_paragraph('Each member must explain a screen during the class demonstration and upload the required artifacts individually.')
for style in doc.styles:
 if style.element.pPr is not None:
  for border in list(style.element.pPr.findall(qn('w:pBdr'))): style.element.pPr.remove(border)
output=root/'submission/Miller_Myles_Activity11_GhostHunt.docx'; doc.save(output); print(output)
